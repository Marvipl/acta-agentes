"""Planilha mestre do plano estratégico (xlsx). Espelha o motor: edite os nós, não a planilha."""
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter as L

N, B = Font(name="Arial", size=10), Font(name="Arial", size=10, bold=True)
AZ = Font(name="Arial", size=10, color="0000FF")
CAB, FUNDO, DEST = Font(name="Arial", size=10, bold=True, color="FFFFFF"), PatternFill("solid", fgColor="111626"), PatternFill("solid", fgColor="F4F4F4")
BRL, PCT = '"R$" #,##0;("R$" #,##0);-', '0.0%;(0.0%);-'


def _cab(ws, cols, larg=None):
    for i, c in enumerate(cols, 1):
        x = ws.cell(row=1, column=i, value=c); x.font, x.fill = CAB, FUNDO
        x.alignment = Alignment(wrap_text=True, vertical="center")
        ws.column_dimensions[L(i)].width = (larg or {}).get(i, 16)
    ws.freeze_panes = "A2"


def _v(ws, r, c, v, fmt=None, font=N):
    x = ws.cell(row=r, column=c, value=v); x.font = font
    if fmt: x.number_format = fmt
    return x


def _dre(ws, titulo, anual, r0):
    anos = list(anual.keys())
    _v(ws, r0, 1, titulo, font=B)
    for j, a in enumerate(anos, 2):
        _v(ws, r0, j, int(a), font=B)
    linhas = [("Receita bruta", "receita", None), ("(-) Impostos sobre receita", "impostos", -1), ("Receita líquida", "=", "r0+1,r0+2"),
              ("(-) CPV", "cpv", -1), ("Lucro bruto", "=", "r0+3,r0+4"), ("(-) Pessoal", "pessoal", -1), ("(-) Outras despesas", "despesas", -1),
              ("EBITDA", "=", "r0+5,r0+6,r0+7"), ("Margem EBITDA", "%", None), ("(+) Fomento e subvenções", "fomento", 1),
              ("(-) Capex", "capex", -1), ("(+) Captação", "captacao", 1), ("Caixa final do ano", "caixa_final", 1)]
    for i, (nome, k, arg) in enumerate(linhas, 1):
        r = r0 + i
        neg = k == "=" or nome.startswith("Caixa")
        _v(ws, r, 1, nome, font=B if neg else N)
        for j, a in enumerate(anos, 2):
            col = L(j)
            if k == "=":
                refs = [eval(x.replace("r0", str(r0))) for x in arg.split(",")]
                _v(ws, r, j, "=" + "+".join(f"{col}{x}" for x in refs), BRL, B)
            elif k == "%":
                _v(ws, r, j, f"=IF({col}{r0+1}=0,0,{col}{r0+8}/{col}{r0+1})", PCT, N)
            else:
                _v(ws, r, j, (arg or 1) * anual[a][k], BRL, AZ)
        if neg:
            for c in range(1, len(anos) + 2): ws.cell(row=r, column=c).fill = DEST
    return r0 + len(linhas) + 2


def gerar(caminho, res, det, nos):
    wb = Workbook()
    ws = wb.active; ws.title = "Resumo"
    _cab(ws, ["Indicador"] + list(res.get("cenarios", {}).keys()), {1: 44, 2: 20, 3: 20, 4: 20})
    inds = [("Menor caixa no horizonte", "menor_caixa", BRL), ("Mês do menor caixa", "mes_menor_caixa", None),
            ("Necessidade de captação (sem captação planejada)", "necessidade_captacao", BRL),
            ("Mês em que o caixa zera sem captação", "mes_caixa_zera_sem_captacao", None),
            ("Primeiro ano com EBITDA positivo", "ano_ebitda_positivo", None), ("Queima média mensal no ano 1", "queima_media_ano1", BRL),
            ("Runway sem captação (meses)", "runway_meses_sem_captacao", None),
            ("Caixa final do horizonte", "caixa_final", BRL)]
    for i, (nome, k, fmt) in enumerate(inds, 2):
        _v(ws, i, 1, nome)
        for j, (c, d) in enumerate(res.get("cenarios", {}).items(), 2):
            _v(ws, i, j, d["indicadores"].get(k), fmt)
    r = len(inds) + 3
    _v(ws, r, 1, "Alertas do motor", font=B)
    for i, a in enumerate(res.get("alertas") or ["nenhum"], r + 1):
        _v(ws, i, 1, a)
    _v(ws, i + 2, 1, "Legenda: azul = valor calculado a partir dos nós; preto = fórmula. Edite os nós e rode o motor.")

    wd = wb.create_sheet("DRE_cenarios"); wd.column_dimensions["A"].width = 36
    r = 1
    for c, d in res.get("cenarios", {}).items():
        r = _dre(wd, f"Cenário {c}", d["anual"], r)
    wo = wb.create_sheet("DRE_opcoes"); wo.column_dimensions["A"].width = 36
    r = 1
    for o, d in res.get("opcoes", {}).items():
        r = _dre(wo, f"Opção {o}", d["anual"], r)

    base = det.get("cenario_base") or []
    wc = wb.create_sheet("Caixa_base")
    cols = ["mes", "receita", "impostos", "cpv", "pessoal", "despesas", "ebitda", "fomento", "capex", "captacao", "caixa"]
    _cab(wc, ["Mês", "Receita", "Impostos", "CPV", "Pessoal", "Outras despesas", "EBITDA", "Fomento", "Capex", "Captação", "Caixa"])
    for i, x in enumerate(base, 2):
        _v(wc, i, 1, x["mes"])
        for j, k in enumerate(cols[1:], 2):
            _v(wc, i, j, x[k], BRL, AZ)

    wp = wb.create_sheet("Portfolio")
    _cab(wp, ["Linha", "Atratividade (0–5)", "Capacidade de vencer (0–5)", "Quadrante", "Decisão"], {1: 34, 4: 40})
    for i, l in enumerate(res.get("portfolio", []), 2):
        for j, k in enumerate(["linha", "atratividade", "capacidade", "quadrante", "decisao"], 1):
            _v(wp, i, j, l[k], "0.00" if j in (2, 3) else None)

    wi = wb.create_sheet("Iniciativas")
    _cab(wi, ["Prioridade", "ID", "Iniciativa", "Objetivo", "Tipo", "Valor", "Confiança", "Esforço (pessoas-mês)", "Início", "Fim", "Custo externo"], {3: 40})
    for i, x in enumerate(res.get("iniciativas", []), 2):
        vals = [x["prioridade"], x["id"], x["nome"], x["objetivo_id"], x["tipo"], x["valor"], x["confianca"], x["esforco_pessoas_mes"], x["inicio"], x["fim"], x["custo_externo"]]
        for j, v in enumerate(vals, 1):
            _v(wi, i, j, v, "0.00" if j == 1 else (BRL if j == 11 else None))
    r = len(res.get("iniciativas", [])) + 3
    _v(wi, r, 1, "Sobrecargas de capacidade (FTE)", font=B)
    for i, s in enumerate(res.get("capacidade", {}).get("sobrecargas") or [{"perfil": "nenhuma"}], r + 1):
        _v(wi, i, 1, s.get("perfil")); _v(wi, i, 2, s.get("mes")); _v(wi, i, 3, s.get("fte_necessario")); _v(wi, i, 4, s.get("fte_disponivel"))

    rm = res.get("roadmap", [])
    tri = sorted({t for r in rm for t in r["trimestres"]})
    wm = wb.create_sheet("Roadmap")
    _cab(wm, ["ID", "Iniciativa", "Área", "Dono", "Prioridade"] + tri, {2: 40, 3: 22, 4: 18})
    for i, r in enumerate(sorted(rm, key=lambda x: (x["area"] or "", -(x["prioridade"] or 0))), 2):
        for j, v in enumerate([r["id"], r["nome"], r["area"], r["dono"], r["prioridade"]], 1):
            _v(wm, i, j, v, "0.00" if j == 5 else None)
        for j, t in enumerate(tri, 6):
            if t in r["trimestres"]:
                c = _v(wm, i, j, "■"); c.fill = PatternFill("solid", fgColor="EE7D00"); c.font = Font(name="Arial", size=10, color="FFFFFF")

    wa = wb.create_sheet("Areas")
    _cab(wa, ["Área", "Dono", "Iniciativas", "Esforço (pessoas-mês)", "Custo externo das iniciativas", "Despesas recorrentes ano 1", "Total"], {1: 26, 2: 18})
    for i, d in enumerate(res.get("areas", []), 2):
        for j, v in enumerate([d["area"], d.get("dono"), d["iniciativas"], d["esforco_pessoas_mes"], d["custo_externo"], d["despesas_recorrentes_ano"]], 1):
            _v(wa, i, j, v, BRL if j in (5, 6) else ("0.0" if j == 4 else None), AZ if j in (5, 6) else N)
        _v(wa, i, 7, f"=E{i}+F{i}", BRL, B)

    wk = wb.create_sheet("OKRs")
    _cab(wk, ["Objetivo", "KR", "Métrica", "Baseline", "Meta anual", "T1", "T2", "T3", "T4", "Dono", "Fonte do dado"], {1: 30, 2: 40})
    ok = nos.get("okrs") or {}
    objs = {o["id"]: o.get("objetivo") for o in ok.get("objetivos", []) or []}
    for i, k in enumerate(ok.get("krs", []) or [], 2):
        mt = k.get("metas_trimestrais") or {}
        vals = [objs.get(k.get("objetivo_id")), k.get("kr"), k.get("metrica"), k.get("baseline"), k.get("meta_anual"),
                mt.get("T1"), mt.get("T2"), mt.get("T3"), mt.get("T4"), k.get("dono"), k.get("fonte_dado")]
        for j, v in enumerate(vals, 1):
            _v(wk, i, j, v)

    wr = wb.create_sheet("Riscos")
    _cab(wr, ["ID", "Risco", "Categoria", "Probabilidade", "Impacto", "Gatilho", "Mitigação", "Plano B", "Dono"], {2: 40, 6: 30, 7: 30, 8: 30})
    for i, x in enumerate((nos.get("riscos") or {}).get("itens", []) or [], 2):
        for j, k in enumerate(["id", "risco", "categoria", "probabilidade", "impacto", "gatilho", "mitigacao", "plano_b", "dono"], 1):
            _v(wr, i, j, x.get(k))

    ww = wb.create_sheet("Hipoteses_WMBT")
    _cab(ww, ["Opção", "ID", "O que precisa ser verdade", "Teste", "Critério de falha", "Prazo", "Status"], {3: 44, 4: 34, 5: 30})
    r = 2
    for o in (nos.get("opcoes") or {}).get("opcoes", []) or []:
        for w in o.get("wmbt", []) or []:
            for j, v in enumerate([o.get("id"), w.get("id"), w.get("condicao"), w.get("teste"), w.get("criterio_de_falha"), w.get("prazo"), w.get("status")], 1):
                _v(ww, r, j, v)
            r += 1
    wb.save(caminho)
