"""Planilha mestre da especificação de produto. Espelha o motor: edite os nós, não a planilha."""
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter as L

N, B, AZ = Font(name="Arial", size=10), Font(name="Arial", size=10, bold=True), Font(name="Arial", size=10, color="0000FF")
CAB, FUNDO, DEST = Font(name="Arial", size=10, bold=True, color="FFFFFF"), PatternFill("solid", fgColor="111626"), PatternFill("solid", fgColor="F4F4F4")
BRL, PCT = '"R$" #,##0;("R$" #,##0);-', '0.0%;(0.0%);-'


def _cab(ws, cols, larg=None):
    for i, c in enumerate(cols, 1):
        x = ws.cell(row=1, column=i, value=c); x.font, x.fill = CAB, FUNDO
        x.alignment = Alignment(wrap_text=True, vertical="center")
        ws.column_dimensions[L(i)].width = (larg or {}).get(i, 16)
    ws.freeze_panes = "A2"


def _v(ws, r, c, v, fmt=None, font=N):
    x = ws.cell(row=r, column=c, value=v); x.font = font
    if fmt: x.number_format = fmt
    return x


def gerar(caminho, res, nos):
    wb = Workbook(); ws = wb.active; ws.title = "Resumo"
    _cab(ws, ["Indicador", "Valor"], {1: 48, 2: 26})
    n, r_ = res.get("negocio") or {}, res.get("roi_cliente") or {}
    itens = [("Nota da oportunidade (0–5)", res["oportunidade"]["nota"], "0.00"), ("Decisão sobre a oportunidade", res["oportunidade"]["decisao"], None),
             ("Conceito escolhido", res["conceitos"]["escolhido"], None), ("Requisitos (total / must)", f"{res['requisitos']['total']} / {res['requisitos']['por_prioridade']['must']}", None),
             ("Custo unitário provável", n.get("custo_unitario"), BRL), ("Margem unitária na venda", n.get("margem_unitaria_venda"), BRL),
             ("VPL do produto em 3 anos", n.get("vpl"), BRL), ("Fluxo acumulado em 3 anos", n.get("acumulado_3anos"), BRL), ("Payback do investimento", n.get("payback") or "além de 3 anos", None),
             ("ROI do cliente: investimento", r_.get("investimento_cliente"), BRL), ("ROI do cliente: economia líquida anual", r_.get("economia_liquida_anual"), BRL),
             ("ROI do cliente: payback (meses)", r_.get("payback_meses"), "0.0")]
    for i, (k, v, f) in enumerate(itens, 2):
        _v(ws, i, 1, k); _v(ws, i, 2, v, f)
    r = len(itens) + 3
    _v(ws, r, 1, "Alertas do motor", font=B)
    for i, a in enumerate(res.get("alertas") or ["nenhum"], r + 1):
        _v(ws, i, 1, a)

    bm = res.get("benchmark") or {}
    wb_ = wb.create_sheet("Benchmark")
    prods = bm.get("produtos", [])
    _cab(wb_, ["Especificação", "Unidade", "Melhor quando"] + [p["nome"] or p["id"] for p in prods] + ["Melhor do mercado", "Líder", "Mediana", "Produtos com dado"], {1: 30})
    nomes = {p["id"]: p["nome"] or p["id"] for p in prods}
    r = 2
    for e in bm.get("especificacoes", []):
        vals = [e["nome"], e["unidade"], e["melhor"]] + [e["valores"].get(p["id"]) for p in prods] + [e["melhor_valor"], nomes.get(e["lider"]), e["mediana"], e["n_dados"]]
        for j, v in enumerate(vals, 1):
            _v(wb_, r, j, v, None, AZ if 3 < j <= 3 + len(prods) else N)
        r += 1
    r += 1
    _v(wb_, r, 1, "Preços", font=B); r += 1
    for j, t in enumerate(["Produto", "Fabricante", "Valor", "Moeda", "Tipo de preço", "Unidade", "Suporte no Brasil"], 1):
        x = _v(wb_, r, j, t, font=B); x.fill = DEST
    r += 1
    for pr in bm.get("precos", []):
        for j, v in enumerate([pr["produto"], pr["fabricante"], pr["valor"], pr["moeda"], pr["tipo"], pr["unidade"], pr["suporte_brasil"]], 1):
            _v(wb_, r, j, v, "#,##0" if j == 3 else None)
        r += 1
    r += 1
    _v(wb_, r, 1, "Requisitos x mercado", font=B); r += 1
    for j, t in enumerate(["Requisito", "Prioridade", "Especificação", "Alvo", "Melhor do mercado", "Líder", "Mediana", "Posição", "Produtos que atendem"], 1):
        x = _v(wb_, r, j, t, font=B); x.fill = DEST
    r += 1
    for a in bm.get("atendimento", []):
        for j, v in enumerate([f"{a['requisito']} · {a['descricao']}", a["prioridade"], f"{a['nome_especificacao']} ({a['unidade'] or ''})", a["alvo"], a["melhor_mercado"], a["lider"], a["mediana"], a["posicao"], ", ".join(a["atendem"]) or "nenhum"], 1):
            _v(wb_, r, j, v)
        r += 1
    r += 1
    _v(wb_, r, 1, f"Cobertura dos dados: {(bm.get('cobertura_dados') or 0)*100:.0f}% das células preenchidas; {(bm.get('dados_confirmados') or 0)*100:.0f}% dos dados com fonte confirmada (ficha técnica ou fabricante).")

    wc = wb.create_sheet("Conceitos")
    _cab(wc, ["Posição", "ID", "Conceito", "Abordagem", "Nota ponderada"], {3: 40})
    for i, c in enumerate(res["conceitos"]["ranking"], 2):
        for j, v in enumerate([i - 1, c["id"], c["nome"], c["abordagem"], c["nota"]], 1):
            _v(wc, i, j, v, "0.00" if j == 5 else None)

    wr = wb.create_sheet("Requisitos")
    _cab(wr, ["ID", "Tipo", "Prioridade", "Descrição", "Origem", "Critério de aceite", "Métrica", "Valor-alvo", "Release"], {4: 46, 6: 40})
    for i, q in enumerate((nos.get("requisitos") or {}).get("itens", []) or [], 2):
        vals = [q.get("id"), q.get("tipo"), q.get("prioridade"), q.get("descricao"), ", ".join(q.get("origem", []) or []), q.get("criterio_aceite"), q.get("metrica"), q.get("valor_alvo"), q.get("release")]
        for j, v in enumerate(vals, 1): _v(wr, i, j, v)

    wa = wb.create_sheet("Arquitetura")
    _cab(wa, ["ID", "Componente", "Tipo", "Decisão", "TRL", "Por unidade", "Custo mín.", "Custo provável", "Custo máx.", "Fonte"], {2: 34, 10: 30})
    for i, c in enumerate((nos.get("arquitetura") or {}).get("componentes", []) or [], 2):
        cu = c.get("custo_unitario") or {}
        vals = [c.get("id"), c.get("nome"), c.get("tipo"), c.get("decisao"), c.get("trl"), c.get("por_unidade"), cu.get("min"), cu.get("provavel"), cu.get("max"), c.get("fonte_custo")]
        for j, v in enumerate(vals, 1): _v(wa, i, j, v, BRL if j in (7, 8, 9) else None, AZ if j in (7, 8, 9) else N)
    k = len((nos.get("arquitetura") or {}).get("componentes", []) or []) + 2
    _v(wa, k, 2, "Custo unitário (componentes por unidade)", font=B)
    _v(wa, k, 8, f'=SUMIF(F2:F{k-1},TRUE,H2:H{k-1})', BRL, B)

    wn = wb.create_sheet("Caso_de_negocio")
    cols = ["unidades_novas", "base_instalada", "receita", "impostos", "comissao", "custo_equipamentos_vendidos", "custo_implantacao", "custo_suporte",
            "margem_contribuicao", "capex_locacao", "investimento_desenvolvimento", "fluxo", "acumulado"]
    _cab(wn, ["Linha"] + [a["ano"] for a in n.get("anos", [])], {1: 36})
    for i, c in enumerate(cols, 2):
        _v(wn, i, 1, c.replace("_", " ").capitalize(), font=B if c in ("margem_contribuicao", "fluxo", "acumulado") else N)
        for j, a in enumerate(n.get("anos", []), 2):
            _v(wn, i, j, a[c], "#,##0" if c in ("unidades_novas", "base_instalada") else BRL, AZ)
    r = len(cols) + 3
    _v(wn, r, 1, "Sensibilidade do VPL (min / max do driver)", font=B)
    for i, s in enumerate(n.get("sensibilidade", []), r + 1):
        _v(wn, i, 1, s["driver"]); _v(wn, i, 2, s["vpl_min"], BRL); _v(wn, i, 3, s["vpl_max"], BRL); _v(wn, i, 4, s["amplitude"], BRL)

    wh = wb.create_sheet("Hipoteses")
    _cab(wh, ["Risco (impacto × incerteza)", "ID", "Hipótese", "Categoria", "Status", "Tem experimento"], {3: 50})
    for i, h in enumerate(res.get("hipoteses", []), 2):
        for j, v in enumerate([h["risco"], h["id"], h["hipotese"], h["categoria"], h["status"], h["tem_experimento"]], 1): _v(wh, i, j, v)

    wm = wb.create_sheet("Roadmap")
    _cab(wm, ["Release", "Objetivo", "Início", "Duração (meses)", "Requisitos", "Marco de decisão"], {2: 40, 5: 30, 6: 36})
    for i, x in enumerate((nos.get("roadmap") or {}).get("releases", []) or [], 2):
        for j, v in enumerate([x.get("id"), x.get("objetivo"), x.get("inicio"), x.get("duracao_meses"), ", ".join(x.get("requisitos", []) or []), x.get("marco_de_decisao")], 1): _v(wm, i, j, v)
    r = len((nos.get("roadmap") or {}).get("releases", []) or []) + 3
    _v(wm, r, 1, "Sobrecargas de capacidade (FTE)", font=B)
    for i, s in enumerate(res.get("capacidade", {}).get("sobrecargas") or [{"perfil": "nenhuma"}], r + 1):
        _v(wm, i, 1, s.get("perfil")); _v(wm, i, 2, s.get("mes")); _v(wm, i, 3, s.get("fte_necessario")); _v(wm, i, 4, s.get("fte_disponivel"))
    wb.save(caminho)
