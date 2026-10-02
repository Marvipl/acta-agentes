"""Importa os documentos do briefing para o orçamento e extrai o texto para os agentes lerem.

Uso (PowerShell, a partir da pasta squad-orcamento):
  python -m motor.insumos importar <dir_versao> <arquivo_ou_pasta> [<arquivo_ou_pasta> ...]
  python -m motor.insumos listar <dir_versao>

O que faz:
  - copia os originais para <dv>/insumos/originais/ (o original nunca é alterado);
  - extrai texto para <dv>/insumos/texto/<arquivo>.md, com marcação de página, aba ou slide para citação;
  - PDF sem camada de texto (digitalizado): gera imagens das páginas em <dv>/insumos/paginas/ para leitura visual;
  - atualiza <dv>/insumos/indice.md e indice.json e a lista nos/briefing.json -> documentos.
Formatos com extração: pdf, docx, xlsx/xlsm, pptx, eml, txt, md, csv, json. Imagens: leitura visual. zip: descompacta e processa.
"""
import argparse, email, hashlib, json, shutil, sys, zipfile
from email import policy
from pathlib import Path
from .util import carregar_json, salvar_json, agora

TEXTO_SIMPLES = {".txt", ".md", ".csv", ".json"}
IMAGENS = {".png", ".jpg", ".jpeg", ".webp", ".gif"}
MAX_PAGINAS_IMAGEM = 30


def _hash(p):
    h = hashlib.sha256()
    with open(p, "rb") as f:
        for bloco in iter(lambda: f.read(1 << 20), b""):
            h.update(bloco)
    return h.hexdigest()[:12]


def _pdf(p, dv):
    try:
        import pypdfium2 as pdfium
    except ImportError:
        return None, "dependência ausente: pip install pypdfium2", 0
    doc = pdfium.PdfDocument(str(p))
    partes, sem_texto = [], 0
    for i in range(len(doc)):
        t = doc[i].get_textpage().get_text_range().strip()
        if len(t) < 20:
            sem_texto += 1
        partes.append(f"\n\n--- página {i + 1} ---\n\n{t}")
    obs = ""
    if sem_texto > len(doc) / 2:
        destino = dv / "insumos" / "paginas" / p.stem
        destino.mkdir(parents=True, exist_ok=True)
        n = min(len(doc), MAX_PAGINAS_IMAGEM)
        for i in range(n):
            pg = doc[i]
            escala = min(2.0, 1400 / max(pg.get_width(), 1))  # até ~1400 px de largura: suficiente para leitura visual
            pg.render(scale=escala).to_pil().convert("RGB").save(destino / f"pagina_{i + 1:03d}.jpg", quality=80)
        obs = f"PDF digitalizado: texto insuficiente; {n} página(s) em insumos/paginas/{p.stem}/ para leitura visual"
        if len(doc) > n:
            obs += f" (só as {n} primeiras de {len(doc)})"
    return "".join(partes), obs, len(doc)


def _docx(p):
    try:
        import docx
    except ImportError:
        return None, "dependência ausente: pip install python-docx", 0
    d = docx.Document(str(p))
    linhas = [par.text for par in d.paragraphs if par.text.strip()]
    for ti, tab in enumerate(d.tables, 1):
        linhas.append(f"\n--- tabela {ti} ---")
        for row in tab.rows:
            linhas.append(" | ".join(c.text.strip() for c in row.cells))
    return "\n".join(linhas), "", 0


MAX_LINHAS, MAX_COLUNAS, MAX_VAZIAS = 5000, 60, 300


def _xlsx(p):
    from openpyxl import load_workbook
    wb = load_workbook(str(p), data_only=True, read_only=True)
    partes, cortes = [], []
    for ws in wb.worksheets:
        partes.append(f"\n\n--- aba {ws.title} ---\n")
        vazias = lidas = 0
        for row in ws.iter_rows(values_only=True, max_col=MAX_COLUNAS):
            lidas += 1
            if any(v is not None and str(v).strip() for v in row):
                vazias = 0
                vals = ["" if v is None else str(v) for v in row]
                while vals and not vals[-1]:
                    vals.pop()
                partes.append(" | ".join(vals))
            else:
                vazias += 1
                if vazias >= MAX_VAZIAS:
                    break
            if lidas >= MAX_LINHAS:
                cortes.append(ws.title); break
    obs = "valores calculados (fórmulas não preservadas no texto)"
    if cortes:
        obs += f"; abas cortadas em {MAX_LINHAS} linhas: {cortes}"
    return "\n".join(partes), obs, len(wb.worksheets)


def _pptx(p):
    try:
        from pptx import Presentation
    except ImportError:
        return None, "dependência ausente: pip install python-pptx", 0
    pr = Presentation(str(p))
    partes = []
    for i, s in enumerate(pr.slides, 1):
        textos = [sh.text_frame.text for sh in s.shapes if getattr(sh, "has_text_frame", False) and sh.text_frame.text.strip()]
        partes.append(f"\n\n--- slide {i} ---\n\n" + "\n".join(textos))
    return "".join(partes), "imagens dos slides não extraídas: abra o original se o conteúdo for visual", len(pr.slides)


def _eml(p):
    with open(p, "rb") as f:
        m = email.message_from_binary_file(f, policy=policy.default)
    corpo = m.get_body(preferencelist=("plain", "html"))
    texto = corpo.get_content() if corpo else ""
    cab = f"De: {m['from']}\nPara: {m['to']}\nData: {m['date']}\nAssunto: {m['subject']}\n\n"
    anexos = [a.get_filename() for a in m.iter_attachments() if a.get_filename()]
    obs = f"anexos do e-mail não extraídos: {anexos} (salve-os e importe)" if anexos else ""
    return cab + texto, obs, 0


def _processar(arq, dv, rel):
    ext = arq.suffix.lower()
    texto, obs, unidades = None, "", 0
    try:
        if ext == ".pdf":
            texto, obs, unidades = _pdf(arq, dv)
        elif ext == ".docx":
            texto, obs, unidades = _docx(arq)
        elif ext in (".xlsx", ".xlsm"):
            texto, obs, unidades = _xlsx(arq)
        elif ext == ".pptx":
            texto, obs, unidades = _pptx(arq)
        elif ext == ".eml":
            texto, obs, unidades = _eml(arq)
        elif ext in TEXTO_SIMPLES:
            texto = arq.read_text(encoding="utf-8", errors="replace")
        elif ext in IMAGENS:
            obs = "imagem: leitura visual do original"
        elif ext in (".doc", ".xls", ".ppt", ".msg"):
            obs = "formato antigo ou do Outlook: salve como .docx, .xlsx, .pptx, .pdf ou .eml e importe de novo"
        else:
            obs = "formato sem extração automática"
    except Exception as e:
        obs = f"erro na extração: {e}"
    destino_txt = None
    if texto:
        destino_txt = dv / "insumos" / "texto" / (rel.replace("/", "__") + ".md")
        destino_txt.parent.mkdir(parents=True, exist_ok=True)
        destino_txt.write_text(f"# {rel}\n\nFonte: insumos/originais/{rel}\n{texto}\n", encoding="utf-8")
    status = "extraido" if texto and not obs.startswith("PDF digitalizado") else ("visual" if ("visual" in obs) else ("parcial" if texto else "nao_extraido"))
    return {"arquivo": rel, "tipo": ext.lstrip("."), "tamanho_kb": round(arq.stat().st_size / 1024, 1), "hash": _hash(arq),
            "paginas_abas_slides": unidades, "texto": str(destino_txt.relative_to(dv)).replace("\\", "/") if destino_txt else "",
            "status": status, "observacao": obs, "importado_em": agora()}


def importar(dv, caminhos):
    dv = Path(dv)
    orig = dv / "insumos" / "originais"
    orig.mkdir(parents=True, exist_ok=True)
    indice = {d["arquivo"]: d for d in carregar_json(dv / "insumos" / "indice.json", [])}
    arquivos = []
    for c in caminhos:
        p = Path(c).expanduser()
        if not p.exists():
            print(f"não encontrado: {p}"); continue
        lista = [p] if p.is_file() else [x for x in sorted(p.rglob("*")) if x.is_file()]
        base = p.parent if p.is_file() else p
        for x in lista:
            rel = str(x.relative_to(base)).replace("\\", "/")
            destino = orig / rel
            destino.parent.mkdir(parents=True, exist_ok=True)
            shutil.copy2(x, destino)
            if x.suffix.lower() == ".zip":
                pasta = destino.with_suffix("")
                with zipfile.ZipFile(destino) as z:
                    z.extractall(pasta)
                for y in sorted(pasta.rglob("*")):
                    if y.is_file():
                        arquivos.append((y, str(y.relative_to(orig)).replace("\\", "/")))
            else:
                arquivos.append((destino, rel))
    for arq, rel in arquivos:
        reg = _processar(arq, dv, rel)
        indice[rel] = reg
        print(f"{reg['status']:13} {rel}  {reg['observacao']}")
    regs = sorted(indice.values(), key=lambda d: d["arquivo"])
    salvar_json(dv / "insumos" / "indice.json", regs)
    linhas = ["# Índice dos insumos do briefing", "", "Leia este índice primeiro. Cite a fonte como `arquivo, página/aba/slide`.", "",
              "| Arquivo | Tipo | Status | Texto extraído | Observação |", "|---|---|---|---|---|"]
    for d in regs:
        linhas.append(f"| {d['arquivo']} | {d['tipo']} | {d['status']} | {d['texto'] or '—'} | {d['observacao'] or ''} |")
    (dv / "insumos" / "indice.md").write_text("\n".join(linhas) + "\n", encoding="utf-8")
    br_p = dv / "nos" / "briefing.json"
    br = carregar_json(br_p, {}) or {}
    br["documentos"] = [{"arquivo": d["arquivo"], "texto": d["texto"], "status": d["status"], "hash": d["hash"]} for d in regs]
    salvar_json(br_p, br)
    print(f"\n{len(arquivos)} arquivo(s) importado(s). Índice: {dv / 'insumos' / 'indice.md'}")
    print("O nó briefing mudou: carimbe de novo (os nós que dependem dele aparecerão como desatualizados).")


def listar(dv):
    for d in carregar_json(Path(dv) / "insumos" / "indice.json", []):
        print(f"{d['status']:13} {d['arquivo']}  {d['observacao']}")


if __name__ == "__main__":
    ap = argparse.ArgumentParser(); sp = ap.add_subparsers(dest="cmd", required=True)
    a = sp.add_parser("importar"); a.add_argument("dv"); a.add_argument("caminhos", nargs="+")
    a = sp.add_parser("listar"); a.add_argument("dv")
    x = ap.parse_args()
    importar(x.dv, x.caminhos) if x.cmd == "importar" else listar(x.dv)
